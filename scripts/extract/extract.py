"""Parse corpus/raw + corpus/quarantine into corpus/parsed (edu-doc-extract conventions).

usage: python -I extract.py PROJECT_ROOT INVENTORY_JSON [--only ID ...] [--types pdf pptx ...] [--fast] [--force]
Restricted records go to corpus/parsed/quarantine/, everything else to corpus/parsed/.
"""
import argparse, ast, io, json, re, sys, time, unicodedata, zipfile
from importlib.metadata import version
from pathlib import Path

ap = argparse.ArgumentParser()
ap.add_argument("root"); ap.add_argument("inventory")
ap.add_argument("--only", nargs="+"); ap.add_argument("--types", nargs="+")
ap.add_argument("--fast", action="store_true", help="PyMuPDF only, no Docling")
ap.add_argument("--no-formula", action="store_true", help="Docling without formula model")
ap.add_argument("--force", action="store_true")
ap.add_argument("--ocr-images", action="store_true", help="pptx: read slide pictures with Docling (formulas + OCR)")
args = ap.parse_args()
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(args.root).resolve(); C = ROOT / "corpus"; OUT = C / "parsed"
recs = [json.loads(l) for l in open(C / "manifest.jsonl", encoding="utf-8")]
inv = json.load(open(args.inventory, encoding="utf-8"))
paths = {}
for x in inv:
    for i in x["ids"]:
        paths.setdefault(i, []).append(x["path"])
DOC_TYPES = {"paper", "slides", "notes", "assignment", "syllabus", "textbook", "notebook", "docs"}
INJECT = re.compile(r"ignore (all |any )?(previous|prior|above) instructions|disregard (the|all) (previous|above)|you are now (an?|the) |system prompt", re.I)
now = lambda: time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())

# ---------------------------------------------------------------- quality
from pylatexenc.latexwalker import LatexWalker

MATH_RE = re.compile(r"\$\$(.+?)\$\$|(?<![\\$])\$([^$\n]+?)\$(?!\$)", re.S)
CODE_RE = re.compile(r"```(\w*)\n(.*?)```", re.S)

def garbage_ratio(t):
    if not t: return 1.0
    bad = sum(1 for ch in t if not (ch.isprintable() or ch in "\n\t") or unicodedata.category(ch) in ("Co", "Cn") or ch == "\ufffd")
    bad += sum(len(m.group()) for m in re.finditer(r"(?:\b\w\b ){8,}", t))  # OCR-ish runs of single letters
    return round(bad / len(t), 4)

def latex_errors(t):
    spans = [a or b for a, b in MATH_RE.findall(t)]
    err = 0
    for s in spans:
        try:
            LatexWalker(s, tolerant_parsing=False).get_latex_nodes()
        except Exception:
            err += 1
    return len(spans), err

def code_errors(t):
    blocks = CODE_RE.findall(t); err = 0
    for lang, body in blocks:
        if lang in ("python", "py") and "# partial" not in body:
            try: ast.parse(body)
            except SyntaxError: err += 1
    return len(blocks), err

def quality(md, n_pages):
    g = garbage_ratio(md); n_eq, le = latex_errors(md); n_code, ce = code_errors(md)
    status = "ok"
    if g > 0.05 or (n_eq and le / n_eq > 0.10): status = "needs_review"
    if len(md.strip()) < 200 and (n_pages or 1) > 1: status = "failed"
    return n_eq, n_code, {"garbage_char_ratio": g, "latex_parse_errors": le, "python_ast_errors": ce, "status": status}

# ---------------------------------------------------------------- parsers
import fitz  # PyMuPDF

def pdf_pymupdf(path):
    doc = fitz.open(path); parts = []; sections = []
    toc = doc.get_toc()
    for lvl, title, page in toc:
        if lvl <= 2: sections.append({"title": title.strip(), "page_start": page})
    for i, page in enumerate(doc, 1):
        parts.append(f"<!-- page: {i} -->\n\n" + page.get_text("text").strip())
    return "\n\n".join(parts), len(doc), sections, f"pymupdf=={version('pymupdf')}"

_conv = {}
def docling_converter(formula):
    if formula not in _conv:
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions, AcceleratorOptions
        from docling.document_converter import DocumentConverter, PdfFormatOption
        o = PdfPipelineOptions(); o.do_ocr = False; o.do_table_structure = True
        o.do_formula_enrichment = formula; o.do_code_enrichment = formula
        o.accelerator_options = AcceleratorOptions(device="auto")
        _conv[formula] = DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=o)})
    return _conv[formula]

def pdf_docling(path, formula):
    res = docling_converter(formula).convert(str(path))
    d = res.document; n = len(d.pages); parts = []
    for p in range(1, n + 1):
        parts.append(f"<!-- page: {p} -->\n\n" + d.export_to_markdown(page_no=p).strip())
    sections = []
    for item, _ in d.iterate_items():
        if getattr(item, "label", None) in ("section_header", "title") and item.prov:
            sections.append({"title": item.text.strip(), "page_start": item.prov[0].page_no})
    tag = f"docling=={version('docling')}" + ("+formula" if formula else "")
    return "\n\n".join(parts), n, sections, tag

def text_pages_ok(path):
    """Selectable text on most pages? (else it is a scan and needs OCR)"""
    doc = fitz.open(path); n = len(doc)
    sample = [doc[i] for i in range(0, n, max(1, n // 8))]
    return sum(len(p.get_text().strip()) > 40 for p in sample) >= len(sample) * 0.6

def pdf(path, kind):
    if not text_pages_ok(path):
        md, n, s, p = pdf_pymupdf(path); return md, n, s, p, ["scanned_or_image_pdf: needs OCR parser"]
    if args.fast:
        md, n, s, p = pdf_pymupdf(path); return md, n, s, p, []
    # slides rarely carry display math worth the formula model's cost; papers/notes/textbook do
    formula = not args.no_formula and kind in ("paper", "notes", "textbook", "assignment")
    try:
        md, n, s, p = pdf_docling(path, formula)
        if garbage_ratio(md) > 0.05:
            raise ValueError("docling garbage ratio high")
        return md, n, s, p, []
    except Exception as e:
        md, n, s, p = pdf_pymupdf(path); return md, n, s, p, [f"docling failed, used fallback: {e.__class__.__name__}: {str(e)[:120]}"]

OMML_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_img_conv = []

def image_converter():
    if not _img_conv:
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
        from docling.document_converter import DocumentConverter, ImageFormatOption
        R = Path.home() / ".cache/docling/models/RapidOcr"  # cached by docling's downloader; modelscope.cn is unreachable
        o = PdfPipelineOptions(); o.do_ocr = True; o.do_formula_enrichment = True; o.do_table_structure = False
        o.ocr_options = RapidOcrOptions(backend="torch", det_model_path=str(R / "PP-OCRv6_det_small.pth"),
                                        rec_model_path=str(R / "PP-OCRv6_rec_small.pth"),
                                        cls_model_path=str(R / "ch_ppocr_mobile_v2.0_cls_mobile.pth"),
                                        rec_keys_path=str(R / "ppocrv6_dict.txt"))
        _img_conv.append(DocumentConverter(allowed_formats=[InputFormat.IMAGE],
                                           format_options={InputFormat.IMAGE: ImageFormatOption(pipeline_options=o)}))
    return _img_conv[0]

def ocr_picture(slide, pic):
    """Render a slide picture (WMF/EMF via Windows GDI, PNG/JPG) and read it with Docling:
    formula regions -> LaTeX, other text -> one OCR'd label line. Returns md lines or None."""
    from PIL import Image
    from docling.datamodel.base_models import DocumentStream
    blip = pic.find(f".//{A_NS}blip")
    if blip is None or not blip.get(R_EMBED): return None
    try:
        blob = slide.part.related_part(blip.get(R_EMBED)).blob
        im = Image.open(io.BytesIO(blob))
        if im.format == "WMF": im.load(dpi=200)
        im = im.convert("RGB")
        if min(im.size) < 40: return None
        buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
        d = image_converter().convert(DocumentStream(name="slide.png", stream=buf)).document
    except Exception:
        return None
    out, words_ = [], []
    for it, _ in d.iterate_items():
        lab = str(getattr(it, "label", "")); txt = (getattr(it, "text", "") or "").strip()
        if not txt: continue
        if lab.endswith("formula"): out.append(f"$$ {txt} $$ <!-- ocr -->")
        else: words_.append(txt)
    if words_: out.append("*Figure text (OCR):* " + " · ".join(words_))
    return out or None

def pptx_parse(path):
    from pptx import Presentation
    from dwml import omml
    prs = Presentation(str(path)); parts = []; sections = []; image_only = []; bad_math = []; n_ocr = [0]
    for i, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title.text.strip() if sl.shapes.title is not None and sl.shapes.title.has_text_frame else ""
        body = []
        for sh in sl.shapes:
            if sh == sl.shapes.title: continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t: body.append("  " * para.level + "- " + t)
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip().replace("|", "\\|").replace("\n", " ") for c in r.cells] for r in sh.table.rows]
                if rows:
                    body.append("| " + " | ".join(rows[0]) + " |\n|" + "---|" * len(rows[0]))
                    body += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        # Office Math (OMML) is invisible to python-pptx runs; convert it, never guess it
        for e in sl._element.iter(OMML_NS + "oMath"):
            try: body.append(f"$$ {omml.oMath2Latex(e).latex.strip()} $$")
            except Exception: body.append("<!-- equation: OMML not convertible -->"); bad_math.append(i)
        pics = list(sl._element.iter(P_NS + "pic"))
        had_text = any(b.startswith(("-", "$$", "|")) for b in body)
        for pic in pics:
            cnv = pic.find(f".//{P_NS}cNvPr")
            alt = (cnv.get("descr") or "").strip() if cnv is not None else ""
            got = ocr_picture(sl, pic) if args.ocr_images else None
            if got:
                body += got; n_ocr[0] += 1
            else:
                body.append(f"*Figure: image not extracted{(' — alt text: ' + alt) if alt else ''}*")
        if pics and not had_text and not any(b.startswith(("$$", "*Figure text")) for b in body):
            image_only.append(i)
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide and sl.notes_slide.notes_text_frame else ""
        block = f"<!-- slide: {i} -->\n\n## {title or f'Slide {i}'}\n\n" + "\n".join(body)
        if notes: block += f"\n\n> Speaker notes: {notes}"
        parts.append(block)
        if title: sections.append({"title": title, "page_start": i})
    msgs = []
    if image_only: msgs.append(f"image-only slides (content possibly equations/figures, not extracted): {image_only}")
    if bad_math: msgs.append(f"unconvertible OMML on slides {bad_math}")
    tag = f"python-pptx=={version('python-pptx')}+dwml"
    if args.ocr_images:
        tag += f"+image-ocr:docling=={version('docling')}(formula+rapidocr)"
        msgs.append(f"{n_ocr[0]} slide images OCR'd; their $$ equations are machine-read (marked ocr) and may contain errors")
    return "\n\n".join(parts), len(prs.slides), sections, tag, msgs

def md_parse(path):
    """Markdown source (d2l.ai). d2l repeats code per framework: keep PyTorch (or untabbed) only,
    drop its build markup (:label:, :begin_tab:, %%tab), normalise fences to ```python."""
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    keep = lambda tab: tab is None or "pytorch" in tab or tab.strip() == "all"
    out, lines, i = [], raw.split("\n"), 0
    while i < len(lines):
        ln = lines[i]
        m = re.match(r":begin_tab:`([^`]*)`", ln.strip())
        if m:
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith(":end_tab:"): j += 1
            if keep(m.group(1)): out += lines[i + 1:j]
            i = j + 1; continue
        if ln.strip().startswith("```{.python"):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"): j += 1
            body = lines[i + 1:j]
            tab = body[0].split("%%tab", 1)[1] if body and body[0].startswith("%%tab") else None
            if keep(tab):
                out += ["```python"] + (body[1:] if tab is not None else body) + ["```"]
            i = j + 1; continue
        if re.match(r"^:(label|eqlabel|width|numref|begin_tab|end_tab)", ln.strip()):
            i += 1; continue
        out.append(ln); i += 1
    md = "\n".join(out)
    md = re.sub(r":(numref|eqref|cite|ref):`([^`]*)`", r"[\2]", md)
    md = re.sub(r"^#\s", "## ", md, flags=re.M)  # the document title comes from the manifest
    sections = [{"title": m.group(1).strip(), "page_start": 1} for m in re.finditer(r"^##+ (.+)$", md, re.M)]
    return md, 1, sections, "d2l-markdown-cleanup+pylatexenc-checked", []

def wiki_math(raw):
    """Wikipedia: replace each rendered math element with its source LaTeX (the alttext), so equations are exact
    rather than lost or OCR'd. Block elements become $$...$$, inline ones $...$."""
    import html as htmllib

    def sub(m):
        block = "mwe-math-element-block" in m.group(1)
        tex = re.search(r'alttext="([^"]*)"', m.group(0))
        if not tex:
            return ""
        t = htmllib.unescape(tex.group(1)).strip()
        t = re.sub(r"^\{\\displaystyle\s*(.*)\}$", r"\1", t, flags=re.S).strip()
        return f"$${t}$$" if block else f"${t}$"
    # the element wraps <math ...alttext=...>...</math> plus a fallback <img>; no nested element of the same class
    return re.sub(r'<span class="(mwe-math-element[^"]*)"[^>]*>.*?</math>.*?(?:<img[^>]*mwe-math-fallback[^>]*>)?\s*</span>',
                  sub, raw, flags=re.S)


def html_parse(path):
    import trafilatura
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    if "wikipedia" in str(path).replace("\\", "/"):
        raw = wiki_math(raw)
    md = trafilatura.extract(raw, output_format="markdown", include_links=True, include_tables=True, include_formatting=True) or ""
    if "wikipedia" in str(path).replace("\\", "/"):
        # undo trafilatura's markdown escaping inside math, drop [edit] links and footnote markers
        md = re.sub(r"\$\$?[^$]+\$\$?", lambda m: re.sub(r"\\([_*\[\]{}#|`~])", r"\1", m.group(0)), md)
        md = re.sub(r"<sup>\\?\[[^\]]{1,12}\\?\]</sup>", "", md)
        md = re.sub(r"^\s*\\?\[edit\\?\]\s*$", "", md, flags=re.M)
        md = re.sub(r"\\\[\s*\[edit\]\([^)]*\)\\\]", "", md)  # "\[ [edit](url)\]" when links are kept
        md = re.sub(r"\[\\\[[^\]]{1,12}\\\]\]\(#cite_note[^)]*\)", "", md)  # "[\[4\]](#cite_note-4)"
        md = re.sub(r"\n{3,}", "\n\n", md)
    sections = [{"title": m.group(2).strip(), "page_start": 1} for m in re.finditer(r"^(#{1,3}) (.+)$", md, re.M)]
    return md, 1, sections, f"trafilatura=={version('trafilatura')}", []

def mdx_parse(path):
    """Hugging Face course .mdx: keep PyTorch branches of {#if fw === 'pt'} blocks, turn quiz <Question> components
    into plain multiple-choice text (correct options marked), keep image alt text as figure captions, and drop the
    remaining JSX components (banners, YouTube embeds, layout divs)."""
    t = Path(path).read_text(encoding="utf-8", errors="replace")
    # framework switch: {#if fw === 'pt'} A {:else} B {/if}  -> A  (and the reverse order -> B)
    t = re.sub(r"\{#if fw === 'pt'\}(.*?)\{:else\}.*?\{/if\}", r"\1", t, flags=re.S)
    t = re.sub(r"\{#if fw === 'tf'\}.*?\{:else\}(.*?)\{/if\}", r"\1", t, flags=re.S)

    def question(m):
        block = m.group(0)
        opts = re.findall(r"text:\s*([\"'`])(.*?)\1\s*,\s*explain:\s*([\"'`])(.*?)\3(\s*,\s*correct:\s*true)?", block, re.S)
        if not opts:
            return ""
        return "\n".join(f"- {'(correct) ' if c else ''}{txt.strip()} — {exp.strip()}" for _, txt, _, exp, c in opts)
    t = re.sub(r"<Question\b.*?/>", question, t, flags=re.S)
    t = re.sub(r"<img\b[^>]*\balt=\"([^\"]+)\"[^>]*/?>", r"*Figure: \1*", t)
    t = re.sub(r"(\*Figure: [^*]+\*)\s*\n\s*\1", r"\1", t)  # light/dark image pairs share one caption
    t = re.sub(r"<([A-Z][A-Za-z]*)\b[^>]*?/>", "", t, flags=re.S)  # self-closing components
    t = re.sub(r"</?(Tip|div|span|details|summary|FrameworkSwitchCourse|DocNotebookDropdown)\b[^>]*>", "", t)
    t = re.sub(r"\[\[[a-z0-9-]+\]\]", "", t)  # heading anchors
    t = re.sub(r"```(\w+) out\b", "```text", t)  # printed outputs are not code
    t = re.sub(r"^#\s", "## ", t, flags=re.M)  # the document title comes from the manifest
    t = re.sub(r"\n{3,}", "\n\n", t).strip()
    sections = [{"title": m.group(1).strip(), "page_start": 1} for m in re.finditer(r"^##+ (.+)$", t, re.M)]
    return t, 1, sections, "hf-course-mdx-cleanup", []


def py_parse(path):
    src = Path(path).read_text(encoding="utf-8", errors="replace")
    sections = [{"title": m.group(1), "page_start": 1} for m in re.finditer(r"^(?:def|class) (\w+)", src, re.M)]
    return f"```python\n{src.rstrip()}\n```", 1, sections, "raw-python", []


def ipynb_parse(path):
    import nbformat
    md = nb_to_md(Path(path).read_bytes())
    md = re.sub(r"^#\s", "## ", md, flags=re.M)
    sections = [{"title": m.group(1).strip(), "page_start": 1} for m in re.finditer(r"^##+ (.+)$", md, re.M)]
    return md, 1, sections, f"nbformat=={version('nbformat')} (outputs stripped)", []


def nb_to_md(raw_bytes):
    import nbformat
    nb = nbformat.reads(raw_bytes.decode("utf-8", "replace"), as_version=4); parts = []
    lang = (nb.metadata.get("kernelspec") or {}).get("language", "python")
    for i, c in enumerate(nb.cells, 1):
        parts.append(f"<!-- cell: {i} type: {c.cell_type} -->")
        parts.append(f"```{lang}\n{c.source}\n```" if c.cell_type == "code" else c.source)
    return "\n\n".join(parts)

CODE_EXT = {".tex": "latex", ".bib": "bibtex", ".py": "python", ".sh": "bash", ".txt": "", ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".cfg": "", ".json": "json"}
def zip_parse(path):
    """Read members in memory only (never extract to disk: untrusted archive)."""
    z = zipfile.ZipFile(path); parts = []; sections = []; notes = []; n = 0
    for m in sorted(z.infolist(), key=lambda m: m.filename):
        name = m.filename
        if m.is_dir() or "__MACOSX" in name or "/." in "/" + name or m.file_size > 5_000_000: continue
        ext = Path(name).suffix.lower(); data = z.read(m)
        if ext == ".ipynb": body = nb_to_md(data)
        elif ext == ".md": body = data.decode("utf-8", "replace")
        elif ext in CODE_EXT: body = f"```{CODE_EXT[ext]}\n{data.decode('utf-8', 'replace')}\n```"
        elif ext == ".pdf":
            d = fitz.open(stream=data, filetype="pdf")
            body = "\n\n".join(f"<!-- page: {i} -->\n\n{p.get_text().strip()}" for i, p in enumerate(d, 1))
        else:
            notes.append(f"skipped member {name}"); continue
        n += 1; parts.append(f"## {name}\n\n{body}"); sections.append({"title": name, "page_start": n})
    return "\n\n".join(parts), n, sections, f"zipfile+nbformat=={version('nbformat')}+pymupdf=={version('pymupdf')}", notes[:30]

# ---------------------------------------------------------------- main
def pick_path(r):
    ps = paths.get(r["id"], [])
    want = "quarantine/" if r["license_tier"] == "restricted" else "raw/"
    pref = [p for p in ps if p.startswith(want)]
    return (pref or ps or [None])[0], len(ps)

def safe_id(i): return re.sub(r"[^A-Za-z0-9._-]+", "_", i)

report = []
for r in recs:
    if r["resource_type"] not in DOC_TYPES: continue
    if args.only and r["id"] not in args.only: continue
    rel, ncopies = pick_path(r)
    if rel is None:
        report.append({"id": r["id"], "status": "failed", "fmt": r["format"], "why": "no file"}); continue
    path = C / rel; ext = path.suffix.lower().lstrip(".")
    if args.types and ext not in args.types: continue
    outdir = OUT / "quarantine" if r["license_tier"] == "restricted" else OUT
    outdir.mkdir(parents=True, exist_ok=True)
    md_path = outdir / f"{safe_id(r['id'])}.md"; js_path = outdir / f"{safe_id(r['id'])}.json"
    if js_path.exists() and not args.force:
        report.append({"id": r["id"], "status": json.load(open(js_path, encoding="utf-8"))["quality"]["status"], "fmt": ext, "skipped": True}); continue
    t0 = time.time()
    try:
        if ext == "pdf": md, n, secs, parser, notes = pdf(path, r["resource_type"])
        elif ext == "pptx": md, n, secs, parser, notes = pptx_parse(path)
        elif ext in ("html", "htm"): md, n, secs, parser, notes = html_parse(path)
        elif ext == "md": md, n, secs, parser, notes = md_parse(path)
        elif ext == "mdx": md, n, secs, parser, notes = mdx_parse(path)
        elif ext == "py": md, n, secs, parser, notes = py_parse(path)
        elif ext == "ipynb": md, n, secs, parser, notes = ipynb_parse(path)
        elif ext == "zip": md, n, secs, parser, notes = zip_parse(path)
        else: raise ValueError(f"no parser for .{ext}")
    except Exception as e:
        report.append({"id": r["id"], "status": "failed", "fmt": ext, "why": f"{e.__class__.__name__}: {e}"[:200]})
        print("FAILED", r["id"], e, flush=True); continue
    md = f"# {r['title']}\n\n{md}\n"
    for i, s in enumerate(secs):  # close section ranges
        s["page_end"] = (secs[i + 1]["page_start"] if i + 1 < len(secs) else n)
    n_eq, n_code, q = quality(md, n)
    side = {"id": r["id"], "title": r["title"], "source_path": rel, "source_sha256": r["sha256"], "parser": parser,
            "parsed_at": now(), "n_pages": n, "n_equations": n_eq, "n_code_blocks": n_code, "sections": secs[:300],
            "quality": q, "injection_suspect": bool(INJECT.search(md)), "notes": notes,
            "topics": r.get("topics"), "license_tier": r["license_tier"], "use": r.get("use"),
            "resource_type": r["resource_type"], "course": r.get("course"), "institution": r.get("institution")}
    if ncopies > 1: side["notes"].append(f"{ncopies} identical copies on disk; parsed {rel}")
    md_path.write_text(md, encoding="utf-8")
    js_path.write_text(json.dumps(side, indent=1, ensure_ascii=False), encoding="utf-8")
    report.append({"id": r["id"], "status": q["status"], "fmt": ext, "parser": parser, "secs": round(time.time() - t0, 1)})
    print(f"{q['status']:12} {ext:5} {r['id']}  ({parser}, {n}p, {n_eq} eq, {round(time.time()-t0)}s)", flush=True)

(OUT / "_reports").mkdir(parents=True, exist_ok=True)
json.dump(report, open(OUT / "_reports" / f"extract-{time.strftime('%Y%m%d-%H%M%S')}.json", "w", encoding="utf-8"), indent=1)
from collections import Counter
print(Counter((x["fmt"], x["status"]) for x in report))
